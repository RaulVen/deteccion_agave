from __future__ import annotations

import csv
from pathlib import Path
from textwrap import shorten

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reportes" / "Reporte_tesis_proceso_completo_agave.docx"
RESULTS = ROOT / "docs" / "results"
RUN = ROOT / "runs" / "detect" / "agave_experiments" / "yolo26s_sgd_cos_slow_120e"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(34, 34, 34)
MUTED = RGBColor(90, 90, 90)
FILL = "F2F4F7"
CALLOUT = "E8EEF5"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": 80, "bottom": 80, "start": 120, "end": 120}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths: list[float], header=True) -> None:
    table.style = "Table Grid"
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(row.cells[i])
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
        for cell in table.rows[0].cells:
            shade(cell, FILL)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Reporte de tesis: proceso completo de deteccion de agave con YOLO26")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = DARK_BLUE
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Desde preparacion del entorno y dataset hasta entrenamiento, seleccion del modelo, inferencia en video y tracking")
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED


def callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = ""
    style_table(table, [6.3], header=False)
    cell = table.rows[0].cells[0]
    shade(cell, CALLOUT)
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(text)
    doc.add_paragraph()


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def kv(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Elemento"
    table.rows[0].cells[1].text = "Detalle"
    for k, v in rows:
        row = table.add_row()
        row.cells[0].text = k
        row.cells[1].text = v
    style_table(table, [1.8, 4.5])


def concepts(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Explicacion para estudiar"
    for k, v in rows:
        row = table.add_row()
        row.cells[0].text = k
        row.cells[1].text = v
    style_table(table, [1.55, 4.75])


def code_block(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    style_table(table, [6.3], header=False)
    cell = table.rows[0].cells[0]
    shade(cell, "F7F7F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(code)
    r.font.name = "Consolas"
    r.font.size = Pt(8)


def image(doc: Document, path: Path, caption: str, width=6.1) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def benchmark_table(doc: Document) -> None:
    rows = []
    with (ROOT / "benchmark" / "agave_benchmark.csv").open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows.append({(k or "").replace("\ufeff", "").strip().strip('"'): (v or "").strip().strip('"') for k, v in row.items()})
    table = doc.add_table(rows=1, cols=8)
    headers = ["Rank", "Run", "Base", "Epocas", "Best", "Prec.", "Recall", "mAP50"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        values = [
            row["rank"],
            row["run"],
            row["model"],
            row["epochs_run"],
            row["best_epoch"],
            row["test_precision"],
            row["test_recall"],
            row["test_mAP50"],
        ]
        for i, value in enumerate(values):
            cells[i].text = value
    style_table(table, [0.42, 1.75, 0.75, 0.55, 0.5, 0.55, 0.6, 0.6])


def architecture_table(doc: Document) -> None:
    rows = [
        ("Backbone", "Conv + C3k2/C3k + SPPF + C2PSA", "Extrae caracteristicas de la imagen: bordes, texturas, formas de hojas y contexto."),
        ("Neck", "Upsample + Concat + C3k2", "Fusiona informacion fina y profunda para detectar agaves de distintos tamanos."),
        ("Head", "Detect", "Genera cajas, confianza y clase agave en multiples escalas."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Parte"
    table.rows[0].cells[1].text = "Bloques"
    table.rows[0].cells[2].text = "Funcion"
    for a, b, c in rows:
        row = table.add_row().cells
        row[0].text = a
        row[1].text = b
        row[2].text = c
    style_table(table, [1.0, 2.2, 3.1])


def layers_table(doc: Document) -> None:
    path = ROOT / "docs" / "model_layers_summary.tsv"
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Bloque"
    table.rows[0].cells[1].text = "Tipo"
    table.rows[0].cells[2].text = "Parametros"
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            idx, layer, params = line.split("\t")
            row = table.add_row().cells
            row[0].text = idx
            row[1].text = layer
            row[2].text = f"{int(params):,}"
    style_table(table, [0.75, 2.25, 1.35])


def read_text(path: str, max_chars: int = 5000) -> str:
    text = (ROOT / path).read_text(encoding="utf-8", errors="replace")
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n...\n[Fragmento recortado para el documento. El archivo completo esta en el repositorio.]"


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    add_title(doc)

    callout(
        doc,
        "Proposito del documento",
        "Este archivo resume todo el proceso realizado para que pueda estudiarse y explicarse en una tesis: entorno, dataset, entrenamientos, metricas, arquitectura, video final y codigo utilizado.",
    )

    doc.add_heading("1. Resumen de lo que se logro", level=1)
    doc.add_paragraph(
        "Se construyo un flujo completo de vision por computadora para detectar agaves. El trabajo inicio con la preparacion del proyecto y del entorno, continuo con la organizacion del dataset, el entrenamiento de multiples modelos YOLO26, la comparacion por metricas, la seleccion del mejor checkpoint y finalmente la aplicacion del modelo sobre video con tracking y conteo."
    )
    bullets(
        doc,
        [
            "Proyecto Git creado en C:\\Users\\iceve\\Documents\\github\\deteccion de agave.",
            "Repositorio remoto configurado en https://github.com/RaulVen/deteccion_agave.git.",
            "Rama principal renombrada a main y subida a GitHub.",
            "Entorno local con Pipenv, Python 3.12.9, Ultralytics, Roboflow, inference-sdk y PyTorch CUDA.",
            "Dataset YOLO26 organizado con 295 imagenes de train, 73 de validation y 36 de test.",
            "Once corridas comparadas en benchmark.",
            "Mejor modelo: yolo26s_sgd_cos_slow_120e.",
            "Video final: segmento de 01:30 a 02:00 con deteccion, tracking, IDs y conteo por linea al 50%.",
            "Resultado del conteo final: 41 agaves unicos cruzando hacia abajo.",
        ],
    )

    doc.add_heading("2. Preparacion del entorno y GPU", level=1)
    doc.add_paragraph(
        "En el proceso no se documento una instalacion manual de driver NVIDIA desde cero dentro del repositorio; lo que se hizo fue preparar el entorno Python y corroborar que la GPU NVIDIA estuviera disponible para entrenamiento. Es decir, se valido que el equipo pudiera usar CUDA desde PyTorch."
    )
    kv(
        doc,
        [
            ("GPU detectada", "NVIDIA GeForce RTX 5050 Laptop GPU, 8 GB"),
            ("Validacion esperada", "nvidia-smi y torch.cuda.is_available() dentro del entorno virtual"),
            ("Python usado", "3.12.9 mediante Pipenv"),
            ("PyTorch", "2.10.0+cu128"),
            ("Torchvision", "0.25.0+cu128"),
            ("Motivo de usar GPU", "Reducir tiempo de entrenamiento y permitir entrenamiento local de YOLO26."),
        ],
    )
    concepts(
        doc,
        [
            ("Driver NVIDIA", "Software del sistema operativo que permite comunicarse con la tarjeta grafica. Sin driver correcto, CUDA no funciona."),
            ("CUDA", "Plataforma de NVIDIA que permite que librerias como PyTorch usen la GPU para operaciones matematicas intensivas."),
            ("PyTorch CUDA", "Version de PyTorch compilada para ejecutar tensores y entrenamiento en GPU."),
            ("nvidia-smi", "Comando que muestra GPU, memoria, version de driver y procesos usando la tarjeta."),
        ],
    )
    doc.add_paragraph("Comandos base usados o equivalentes para validar entorno:")
    code_block(
        doc,
        "$env:PIPENV_VENV_IN_PROJECT='1'\npy -3.12 -m pipenv install\npy -3.12 -m pipenv run python -c \"import torch; print(torch.cuda.is_available()); print(torch.cuda.get_device_name(0))\"\npy -3.12 -m pipenv run nvidia-smi",
    )

    doc.add_heading("3. Dependencias instaladas", level=1)
    doc.add_paragraph(
        "Se uso Pipenv en lugar de Conda. Esto dejo registradas las dependencias en Pipfile y Pipfile.lock, lo cual ayuda a reproducir el entorno."
    )
    code_block(doc, read_text("Pipfile", max_chars=1800))

    doc.add_heading("4. Dataset y organizacion YOLO", level=1)
    doc.add_paragraph(
        "El dataset se organizo en formato YOLO26. En este formato, las imagenes se separan por carpetas y cada imagen tiene un archivo de etiqueta asociado con coordenadas normalizadas de las cajas. La clase del proyecto fue una sola: agave."
    )
    kv(
        doc,
        [
            ("Ruta del dataset", "datasets/yolo26"),
            ("Train", "295 imagenes para aprendizaje"),
            ("Validation", "73 imagenes para medir desempeno durante entrenamiento"),
            ("Test", "36 imagenes reservadas para evaluacion final"),
            ("Clase", "agave"),
            ("Archivo principal", "datasets/yolo26/data.yaml"),
        ],
    )
    concepts(
        doc,
        [
            ("Train", "Datos que el modelo ve durante el entrenamiento."),
            ("Validation", "Datos para revisar si el modelo mejora sin tocar el test final."),
            ("Test", "Datos que sirven para evaluar el modelo ya entrenado."),
            ("Etiqueta", "Archivo con la ubicacion de cada agave en la imagen."),
            ("Bounding box", "Caja rectangular que encierra al objeto detectado."),
        ],
    )
    code_block(doc, read_text("datasets/yolo26/data.yaml", max_chars=800))

    doc.add_heading("5. Estrategia de entrenamiento", level=1)
    doc.add_paragraph(
        "La estrategia fue iterativa. Primero se hizo una prueba rapida para comprobar que el pipeline funcionara. Despues se ajustaron aumentaciones, epocas, optimizadores, learning rates y pesos de perdida. Finalmente se compararon las corridas en un benchmark."
    )
    numbered(
        doc,
        [
            "Validar que YOLO pudiera leer el dataset y entrenar con GPU.",
            "Probar yolo26n como modelo pequeno y rapido.",
            "Probar yolo26s como modelo con mayor capacidad.",
            "Modificar data augmentation segun condiciones esperadas en agave.",
            "Quitar mosaico moderado, rotacion de 10 grados y translacion cuando se observo que podian no ayudar.",
            "Agregar flip vertical y conservar flip horizontal.",
            "Explorar learning rate, optimizador y pesos de loss.",
            "Comparar metricas finales en test.",
        ],
    )

    doc.add_heading("6. Conceptos de entrenamiento", level=1)
    concepts(
        doc,
        [
            ("Epoca", "Una pasada completa por el conjunto de entrenamiento."),
            ("Batch", "Cantidad de imagenes procesadas antes de actualizar pesos. Se uso batch=8 por memoria GPU."),
            ("Learning rate", "Tamano del paso con el que el modelo aprende. Si es alto puede ser inestable; si es bajo aprende lento."),
            ("Optimizador", "Metodo matematico que actualiza los pesos del modelo. Se probaron SGD y AdamW."),
            ("Loss", "Error interno que el modelo intenta reducir durante entrenamiento."),
            ("box loss", "Error relacionado con ubicacion y tamano de cajas."),
            ("cls loss", "Error de clasificacion de la clase agave."),
            ("dfl loss", "Ayuda a refinar la localizacion de las cajas."),
            ("Early stopping", "Detiene el entrenamiento si ya no mejora, evitando sobreentrenamiento."),
            ("Data augmentation", "Transformaciones aplicadas a imagenes para simular variedad."),
        ],
    )

    doc.add_heading("7. Configuracion ganadora", level=1)
    doc.add_paragraph(
        "El mejor resultado se obtuvo con yolo26s_sgd_cos_slow_120e. Aunque se configuraron 120 epocas, el entrenamiento corrio 86 y el mejor checkpoint se encontro en la epoca 56 segun el benchmark final."
    )
    kv(
        doc,
        [
            ("Modelo base", "yolo26s.pt"),
            ("Optimizador", "SGD"),
            ("Learning rate inicial", "0.005"),
            ("Scheduler", "cos_lr=true"),
            ("Patience", "30"),
            ("Batch", "8"),
            ("imgsz", "640"),
            ("Mosaic", "0.0"),
            ("Rotacion", "0 grados"),
            ("Translacion", "0.0"),
            ("Flip vertical", "0.5"),
            ("Flip horizontal", "0.5"),
            ("MixUp", "0.05"),
            ("Pesos de perdida", "box=7.5, cls=0.5, dfl=1.5"),
        ],
    )
    code_block(doc, read_text("configs/training/historico_2026-05_06/yolo26s_sgd_cos_slow_120e.yaml", max_chars=4500))

    doc.add_heading("8. Todos los entrenamientos comparados", level=1)
    doc.add_paragraph(
        "Se generaron 11 corridas y se ordenaron por metricas de test. La tabla siguiente resume el benchmark utilizado para decidir el mejor modelo."
    )
    benchmark_table(doc)
    callout(
        doc,
        "Modelo final seleccionado",
        "El modelo yolo26s_sgd_cos_slow_120e quedo como mejor version general con precision 0.693, recall 0.615, mAP50 0.654 y mAP50-95 0.213 en test.",
    )

    doc.add_heading("9. Metricas para interpretar el resultado", level=1)
    concepts(
        doc,
        [
            ("Precision", "De todas las detecciones que hizo el modelo, cuantas fueron correctas. Alta precision significa pocos falsos positivos."),
            ("Recall", "De todos los agaves reales, cuantos encontro. Alto recall significa pocos agaves perdidos."),
            ("mAP50", "Promedio de precision cuando la caja predicha se solapa al menos 50% con la caja real."),
            ("mAP50-95", "Metrica mas estricta que evalua solapes desde 50% hasta 95%."),
            ("ACC operacional", "Metrica agregada en reportes visuales basada en F1 para explicar balance entre precision y recall. No sustituye mAP."),
        ],
    )

    doc.add_heading("10. Evidencia visual del entrenamiento", level=1)
    image(doc, RESULTS / "best_yolo26s_results.png", "Curvas de entrenamiento del mejor modelo.", width=6.1)
    image(doc, RESULTS / "best_yolo26s_operational_accuracy.png", "Grafica de ACC operacional junto con precision, recall y mAP50.", width=6.1)
    image(doc, RESULTS / "best_yolo26s_confusion_matrix_readable.png", "Matriz de confusion redibujada con contraste alto.", width=5.5)
    image(doc, RESULTS / "best_yolo26s_val_predictions.jpg", "Predicciones del mejor modelo sobre datos de validacion.", width=6.1)

    doc.add_heading("11. Arquitectura del modelo a grandes rasgos", level=1)
    doc.add_paragraph(
        "El checkpoint final corresponde a un modelo de deteccion tipo YOLO con 24 bloques principales y 9,948,638 parametros. A grandes rasgos, la red se divide en backbone, neck y head."
    )
    architecture_table(doc)
    concepts(
        doc,
        [
            ("Conv2D", "Capa convolucional que aplica filtros sobre la imagen para detectar patrones locales como bordes, lineas y texturas."),
            ("BatchNorm", "Normaliza activaciones para estabilizar y acelerar entrenamiento."),
            ("SiLU", "Funcion de activacion que introduce no linealidad para aprender relaciones complejas."),
            ("C3k2/C3k", "Bloques compuestos que reutilizan informacion y mejoran aprendizaje sin aumentar demasiado el costo."),
            ("SPPF", "Pooling piramidal rapido que permite capturar contexto a multiples escalas."),
            ("C2PSA", "Bloque de atencion que ayuda a enfatizar regiones o canales relevantes."),
            ("Upsample", "Aumenta resolucion de mapas para recuperar detalle espacial."),
            ("Concat", "Fusiona caracteristicas de distintas capas."),
            ("Detect", "Capa final que produce cajas, confianza y clase agave."),
        ],
    )
    layers_table(doc)

    doc.add_heading("12. Como se paso del modelo al video final", level=1)
    doc.add_paragraph(
        "Despues de seleccionar el mejor modelo, se aplico inferencia al video 3.mp4. Se recorto el segmento solicitado del minuto 1:30 al 2:00 y se genero una prediccion visual. Luego se agrego tracking para asignar IDs y conteo por linea horizontal."
    )
    numbered(
        doc,
        [
            "Se tomo el video fuente C:\\Users\\iceve\\Documents\\3.mp4.",
            "Se extrajo el segmento de 01:30 a 02:00.",
            "Se aplico el modelo best.pt de yolo26s_sgd_cos_slow_120e.",
            "Se uso conf=0.30 e iou=0.40 para reducir duplicados y controlar confianza.",
            "Se uso tracking BoT-SORT para asignar IDs a los agaves.",
            "Se coloco una linea horizontal al 50% del frame.",
            "Se conto solamente la direccion hacia abajo.",
            "Se corrigio la logica para contar cada ID una sola vez usando --unique-id-only.",
            "Resultado final: 41 IDs unicos contados cruzando hacia abajo.",
        ],
    )
    kv(
        doc,
        [
            ("Video final", "video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.mp4"),
            ("CSV final", "video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.csv"),
            ("Modelo usado", "runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt"),
            ("Linea", "50% del frame"),
            ("Direccion", "down / hacia abajo"),
            ("Conteo", "41 agaves unicos"),
        ],
    )

    doc.add_heading("13. Tracking y conteo", level=1)
    doc.add_paragraph(
        "El tracking es necesario porque en video no basta con detectar objetos en cada frame. Si se cuenta cada deteccion individual, el mismo agave podria sumarse muchas veces. El seguimiento asigna un ID para reconocer el mismo objeto a traves del tiempo."
    )
    concepts(
        doc,
        [
            ("BoT-SORT", "Algoritmo de tracking usado por Ultralytics para asociar detecciones entre frames."),
            ("track_buffer", "Cantidad de frames que se conserva un track aunque no se detecte momentaneamente."),
            ("ID unico", "Identificador del objeto seguido. Permite contar cada agave una vez."),
            ("Cruce de linea", "Evento que ocurre cuando el centro del objeto pasa de un lado de la linea al otro."),
            ("duplicate suppression", "Logica para evitar contar dos veces el mismo evento cuando los IDs cambian o se pierde la deteccion."),
        ],
    )
    code_block(doc, read_text("configs/tracking/botsort_agave.yaml", max_chars=1200))

    doc.add_heading("14. Codigo principal utilizado", level=1)
    doc.add_paragraph(
        "A continuacion se incluyen comandos y fragmentos clave. El script completo esta en el repositorio, pero aqui se deja lo esencial para entender el flujo."
    )
    doc.add_heading("14.1 Comando de tracking final", level=2)
    code_block(
        doc,
        "$env:PIPENV_VENV_IN_PROJECT='1'\npy -3.12 -m pipenv run python scripts/track_count_line.py `\n  --model \"runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt\" `\n  --source \"video_predictions/3_segment_01m30s_02m00s.mp4\" `\n  --output \"video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.mp4\" `\n  --csv \"video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.csv\" `\n  --tracker \"configs/tracking/botsort_agave.yaml\" `\n  --line-ratio 0.50 `\n  --direction down `\n  --conf 0.30 `\n  --iou 0.40 `\n  --imgsz 640 `\n  --device 0 `\n  --unique-id-only `\n  --show-conf",
    )
    doc.add_heading("14.2 Fragmento del script de conteo", level=2)
    code_block(doc, read_text("scripts/track_count_line.py", max_chars=6200))

    doc.add_heading("15. Limitaciones y aprendizaje para tesis", level=1)
    bullets(
        doc,
        [
            "El dataset todavia es pequeno para generalizar a muchas condiciones de campo.",
            "mAP50-95 bajo indica que las cajas pueden mejorar con mas datos y etiquetas mas precisas.",
            "El modelo ya funciona como primera version para detectar y contar en un segmento de video.",
            "El conteo debe validarse contra conteo manual para reportar error real.",
            "Si en videos largos se pierden IDs, el siguiente paso seria integrar DeepSORT o ReID.",
            "Para tesis, se puede presentar como pipeline experimental reproducible: dataset, entrenamiento, evaluacion e inferencia aplicada.",
        ],
    )

    doc.add_heading("16. Como explicarlo en una defensa o presentacion", level=1)
    numbered(
        doc,
        [
            "Primero explicar el problema: detectar agaves automaticamente en imagen/video.",
            "Despues explicar los datos: 404 imagenes, una clase, formato YOLO.",
            "Luego explicar el entorno: GPU NVIDIA, PyTorch CUDA y Pipenv.",
            "Presentar la estrategia: multiples modelos y parametros, comparados por benchmark.",
            "Mostrar el mejor modelo y sus metricas.",
            "Explicar arquitectura: backbone extrae rasgos, neck fusiona escalas, head predice cajas.",
            "Mostrar el video final: deteccion + ID + linea + conteo.",
            "Cerrar con resultado: 41 agaves unicos contados en el segmento y siguientes mejoras.",
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
