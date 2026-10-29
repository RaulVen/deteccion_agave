from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reportes" / "Reporte_avances_deteccion_agave.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(34, 34, 34)
MUTED = RGBColor(90, 90, 90)
HEADER_FILL = "F2F4F7"
CALLOUT_FILL = "E8EEF5"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, widths_inches: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_inches):
            cell = row.cells[idx]
            cell.width = Inches(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, widths_inches: list[float], header=True) -> None:
    table.style = "Table Grid"
    set_table_width(table, widths_inches)
    if header:
        set_repeat_table_header(table.rows[0])
        for cell in table.rows[0].cells:
            set_cell_shading(cell, HEADER_FILL)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9)
    for row in table.rows[1:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def add_callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    style_table(table, [6.3], header=False)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    title_run = p.add_run(title + ": ")
    title_run.bold = True
    title_run.font.color.rgb = DARK_BLUE
    body_run = p.add_run(text)
    body_run.font.color.rgb = INK
    doc.add_paragraph()


def add_key_value_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Elemento"
    table.rows[0].cells[1].text = "Detalle"
    for key, value in rows:
        row = table.add_row()
        row.cells[0].text = key
        row.cells[1].text = value
    style_table(table, [1.85, 4.45])


def add_concept_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Explicacion sencilla"
    for concept, explanation in rows:
        row = table.add_row()
        row.cells[0].text = concept
        row.cells[1].text = explanation
    style_table(table, [1.65, 4.65])


def add_benchmark_table(doc: Document) -> None:
    rows = [
        ("1", "yolo26s_sgd_cos_slow_120e", "yolo26s.pt", "86", "56", "0.693", "0.615", "0.654", "0.213"),
        ("2", "yolo26n_field_aug_60e", "yolo26n.pt", "60", "55", "0.637", "0.539", "0.596", "0.201"),
        ("3", "yolo26s_adamw_box_focus_120e", "yolo26s.pt", "105", "75", "0.659", "0.509", "0.581", "0.200"),
        ("4", "yolo26n_medium_aug_30e", "yolo26n.pt", "30", "28", "0.614", "0.558", "0.598", "0.197"),
        ("5", "yolo26n_field_aug_90e", "yolo26n.pt", "79", "54", "0.631", "0.539", "0.599", "0.192"),
        ("6", "yolo26s_adamw_cos_low_lr_120e", "yolo26s.pt", "57", "27", "0.633", "0.521", "0.589", "0.192"),
        ("7", "yolo26s_field_aug_60e", "yolo26s.pt", "60", "40", "0.622", "0.583", "0.609", "0.192"),
        ("8", "yolo26s_loss_box_dfl_45e", "yolo26s.pt", "45", "40", "0.617", "0.529", "0.571", "0.186"),
        ("9", "yolo26s_coslr_balanced_45e", "yolo26s.pt", "44", "32", "0.549", "0.603", "0.567", "0.183"),
        ("10", "yolo26s_loss_precision_45e", "yolo26s.pt", "20", "8", "0.596", "0.546", "0.526", "0.157"),
        ("11", "agave_quick_test", "yolo26n.pt", "3", "3", "0.431", "0.407", "0.370", "0.109"),
    ]
    headers = ["Rank", "Corrida", "Base", "Ep.", "Best", "Prec.", "Recall", "mAP50", "mAP50-95"]
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for data in rows:
        row = table.add_row()
        for idx, value in enumerate(data):
            row.cells[idx].text = value
    style_table(table, [0.42, 1.72, 0.78, 0.45, 0.47, 0.54, 0.56, 0.56, 0.72])
    for row in table.rows[1:]:
        for idx in [0, 3, 4, 5, 6, 7, 8]:
            row.cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_image_if_exists(doc: Document, image_path: Path, caption: str, width=6.2) -> None:
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Deteccion de Agave | Reporte de avances")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("Reporte ampliado de avances: Deteccion de Agave con YOLO26")
    run.bold = True
    run.font.size = Pt(21)
    run.font.color.rgb = DARK_BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Entrenamiento, evaluacion, inferencia en video, tracking y conteo de objetos")
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Proyecto local: C:\\Users\\iceve\\Documents\\github\\deteccion de agave")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_title(doc)

    add_callout(
        doc,
        "Resumen ejecutivo",
        "Se desarrollo un flujo completo para deteccion automatica de agave: organizacion del dataset, entorno local con GPU, "
        "entrenamientos comparativos de YOLO26, seleccion del mejor modelo, inferencia sobre video y una simulacion de conteo mediante "
        "tracking con IDs unicos. El resultado mas importante de la ultima etapa fue contar 41 agaves cruzando una linea horizontal al "
        "50% del frame en el segmento de video analizado.",
    )

    doc.add_heading("1. Objetivo general", level=1)
    doc.add_paragraph(
        "El objetivo principal fue construir una primera version funcional de un sistema de vision por computadora capaz de detectar "
        "agaves en imagenes y videos. La idea no fue solamente entrenar un modelo, sino crear un flujo completo: preparar datos, entrenar, "
        "comparar resultados, probar en video y comenzar a medir conteos con seguimiento de objetos."
    )
    doc.add_paragraph(
        "Este avance sirve como base para una aplicacion futura de monitoreo agricola, donde el sistema pueda ayudar a identificar plantas "
        "en campo, estimar presencia o movimiento dentro de una zona de interes y apoyar tareas de conteo o inspeccion visual."
    )
    add_bullets(
        doc,
        [
            "Se creo el proyecto Git y se conecto al repositorio remoto de GitHub.",
            "Se preparo un dataset YOLO con 404 imagenes distribuidas en train, validation y test.",
            "Se instalo un entorno Pipenv con librerias para YOLO, Roboflow, inferencia y entrenamiento con GPU.",
            "Se entrenaron y compararon 11 configuraciones diferentes.",
            "Se selecciono el mejor modelo segun metricas de prueba.",
            "Se realizo una simulacion de video con deteccion, tracking de IDs y conteo por cruce de linea.",
        ],
    )

    doc.add_heading("2. Conceptos base para entender el proyecto", level=1)
    doc.add_paragraph(
        "Esta seccion resume los conceptos principales para que el reporte pueda presentarse tambien ante personas que no conocen "
        "entrenamiento de modelos de vision por computadora."
    )
    add_concept_table(
        doc,
        [
            ("Vision por computadora", "Area de inteligencia artificial que permite que una computadora analice imagenes o videos para reconocer patrones, objetos o eventos."),
            ("Deteccion de objetos", "Tarea donde el modelo no solo dice si existe un objeto, sino que dibuja una caja alrededor de el y lo clasifica."),
            ("YOLO", "Familia de modelos de deteccion que procesa una imagen y predice objetos de forma rapida. YOLO significa You Only Look Once."),
            ("YOLO26", "Version usada en el proyecto como base para entrenar modelos de deteccion de agave."),
            ("Modelo preentrenado", "Modelo que ya aprendio caracteristicas generales en otros datos y se ajusta despues al caso especifico, en este caso agaves."),
            ("Dataset", "Conjunto de imagenes y etiquetas que se usa para entrenar y evaluar el modelo."),
            ("Etiqueta o label", "Archivo que indica donde esta cada agave dentro de una imagen, normalmente mediante coordenadas de una caja."),
            ("Bounding box", "Caja rectangular que rodea el objeto detectado."),
            ("Inferencia", "Uso del modelo ya entrenado para hacer predicciones sobre imagenes o videos nuevos."),
            ("Tracking", "Seguimiento de un objeto a traves de varios frames de un video, asignandole un ID para reconocerlo en el tiempo."),
            ("Conteo por linea", "Metodo donde se cuenta un objeto cuando su trayectoria cruza una linea definida en el frame."),
        ],
    )

    doc.add_heading("3. Repositorio y versionamiento", level=1)
    doc.add_paragraph(
        "El proyecto se organizo como repositorio Git para conservar scripts, configuraciones y documentacion. Los archivos pesados, como "
        "pesos entrenados, datasets completos, videos y carpetas de resultados, se dejaron fuera del versionamiento con .gitignore."
    )
    add_key_value_table(
        doc,
        [
            ("Repositorio local", r"C:\Users\iceve\Documents\github\deteccion de agave"),
            ("Remoto GitHub", "https://github.com/RaulVen/deteccion_agave.git"),
            ("Rama actual", "main"),
            ("Commit base documentado", "8ef03d4 - Document agave YOLO experiments and tracking setup"),
            ("Razon del .gitignore", "Evitar subir a GitHub archivos muy grandes o generados automaticamente."),
        ],
    )

    doc.add_heading("4. Dataset y preparacion de datos", level=1)
    doc.add_paragraph(
        "El dataset se obtuvo desde Roboflow y se organizo localmente en formato YOLO26. En este formato, cada imagen tiene un archivo "
        "asociado con las coordenadas de los objetos marcados. Para este avance solo se trabajo con una clase: agave."
    )
    add_key_value_table(
        doc,
        [
            ("Formato", "YOLO26 para deteccion de objetos"),
            ("Clase", "agave"),
            ("Imagenes totales", "404"),
            ("Train", "295 imagenes para que el modelo aprenda"),
            ("Validation", "73 imagenes para monitorear desempeno durante entrenamiento"),
            ("Test", "36 imagenes para evaluar el resultado final"),
            ("Archivo de datos", "datasets/yolo26/data.yaml"),
        ],
    )
    add_concept_table(
        doc,
        [
            ("Train", "Parte del dataset que el modelo ve durante el entrenamiento. Con estas imagenes aprende a detectar agaves."),
            ("Validation", "Parte usada durante el entrenamiento para medir si el modelo esta aprendiendo bien o si empieza a memorizar."),
            ("Test", "Parte reservada para evaluar el modelo al final. Es importante porque simula imagenes no vistas por el modelo."),
            ("Data leakage", "Problema que ocurre si imagenes muy parecidas aparecen en train y test. Puede hacer que las metricas se vean mejores de lo real."),
            ("Calidad de etiquetas", "Mientras mejor esten dibujadas las cajas del agave, mejor puede aprender el modelo. Etiquetas inconsistentes afectan el resultado."),
        ],
    )

    doc.add_heading("5. Entorno local y GPU", level=1)
    doc.add_paragraph(
        "Se preparo el entorno con Pipenv, no Conda, para que las dependencias del proyecto quedaran controladas. La parte critica fue "
        "instalar PyTorch con soporte CUDA, porque el entrenamiento de modelos de vision es mucho mas rapido cuando se usa GPU."
    )
    add_key_value_table(
        doc,
        [
            ("Python", "3.12.9 dentro de Pipenv"),
            ("GPU", "NVIDIA GeForce RTX 5050 Laptop GPU, 8 GB"),
            ("PyTorch", "2.10.0+cu128"),
            ("Torchvision", "0.25.0+cu128"),
            ("Framework principal", "Ultralytics / YOLO26"),
            ("Servicios previstos", "Roboflow e inference-sdk"),
        ],
    )
    add_concept_table(
        doc,
        [
            ("GPU", "Procesador especializado para calculos paralelos. En entrenamiento de redes neuronales acelera mucho el proceso."),
            ("CUDA", "Tecnologia de NVIDIA que permite que PyTorch use la GPU para entrenar modelos."),
            ("Pipenv", "Herramienta para crear un entorno virtual y registrar librerias en Pipfile y Pipfile.lock."),
            ("Batch", "Cantidad de imagenes que procesa el modelo antes de actualizar sus pesos. En este proyecto se uso batch=8 por memoria GPU."),
            ("imgsz", "Tamano al que se redimensionan las imagenes durante entrenamiento e inferencia. Se uso 640."),
        ],
    )

    doc.add_heading("6. Como se entrenaron los modelos", level=1)
    doc.add_paragraph(
        "El entrenamiento consistio en ajustar modelos YOLO26 preentrenados para que aprendieran especificamente la clase agave. "
        "Se probaron modelos base yolo26n y yolo26s. La version n es mas pequena y rapida; la version s tiene mas capacidad y normalmente "
        "puede aprender patrones mas complejos, aunque consume mas recursos."
    )
    add_numbered(
        doc,
        [
            "Primero se hizo una prueba rapida para validar que dataset, entorno, GPU y scripts funcionaran.",
            "Despues se entrenaron configuraciones con mas epocas para observar si el modelo mejoraba.",
            "Se probaron cambios de data augmentation para representar variaciones reales del campo.",
            "Se retiro mosaico moderado, rotacion de 10 grados y translacion, porque podian alterar demasiado la apariencia real del agave.",
            "Se agrego flip vertical y se mantuvo flip horizontal para explorar variaciones visuales sin desplazar las cajas.",
            "Se probaron optimizadores, learning rates y pesos de funcion de perdida para estudiar estabilidad y precision.",
        ],
    )
    add_concept_table(
        doc,
        [
            ("Epoca", "Una pasada completa por las imagenes de entrenamiento. Mas epocas permiten aprender mas, pero tambien pueden causar sobreajuste."),
            ("Early stopping", "Mecanismo que detiene el entrenamiento cuando el modelo deja de mejorar. Evita gastar tiempo y reducir generalizacion."),
            ("Learning rate", "Tamano de los pasos que da el modelo al aprender. Si es muy alto puede inestabilizarse; si es muy bajo puede aprender lento."),
            ("Optimizador", "Algoritmo que decide como actualizar los pesos del modelo. Se probaron SGD y AdamW."),
            ("Funcion de perdida", "Medida interna del error durante entrenamiento. El modelo intenta reducirla."),
            ("box loss", "Parte de la perdida relacionada con que tan bien estan ubicadas las cajas."),
            ("cls loss", "Parte de la perdida relacionada con clasificar correctamente el objeto."),
            ("DFL", "Perdida que ayuda a mejorar la precision de la localizacion de las cajas."),
        ],
    )

    doc.add_heading("7. Aumentacion de datos", level=1)
    doc.add_paragraph(
        "La aumentacion de datos consiste en transformar imagenes durante entrenamiento para que el modelo vea mas variedad sin tener que "
        "capturar nuevas imagenes. Sin embargo, no toda aumentacion ayuda: si transforma demasiado la planta o la escena, puede ensenar "
        "ejemplos poco realistas."
    )
    add_key_value_table(
        doc,
        [
            ("Mosaic", "Desactivado en la mejor corrida: 0.0"),
            ("Rotacion", "Eliminada: 0 grados"),
            ("Translacion", "Eliminada: 0.0"),
            ("Scale", "0.35 para variar tamano del objeto"),
            ("Shear", "2 para leve deformacion geometrica"),
            ("Perspective", "0.0005 para cambio minimo de perspectiva"),
            ("Flip vertical", "0.5"),
            ("Flip horizontal", "0.5"),
            ("MixUp", "0.05"),
            ("Erasing", "0.15"),
        ],
    )
    add_callout(
        doc,
        "Decision importante",
        "No se aumento imagenes de forma agresiva. La recomendacion fue no depender solo de aumentacion sintetica; para mejorar realmente "
        "el modelo conviene agregar mas imagenes reales de agaves en distintas condiciones.",
    )

    doc.add_heading("8. Configuracion ganadora", level=1)
    doc.add_paragraph(
        "La mejor corrida fue yolo26s_sgd_cos_slow_120e. Se programaron 120 epocas, pero el entrenamiento se detuvo a las 86 epocas por "
        "early stopping. El mejor checkpoint se obtuvo en la epoca 56, lo cual significa que despues de esa etapa el modelo ya no mejoro "
        "suficientemente en validacion."
    )
    add_key_value_table(
        doc,
        [
            ("Modelo base", "yolo26s.pt"),
            ("Nombre de corrida", "yolo26s_sgd_cos_slow_120e"),
            ("Optimizador", "SGD"),
            ("Learning rate inicial", "0.005"),
            ("Scheduler", "Cosine LR activado"),
            ("Epocas configuradas", "120"),
            ("Epocas corridas", "86"),
            ("Mejor epoca", "56"),
            ("Patience", "30"),
            ("Batch", "8"),
            ("Image size", "640"),
            ("Device", "0 / GPU"),
            ("Pesos de perdida", "box=7.5, cls=0.5, dfl=1.5"),
            ("Ruta del mejor peso", r"runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt"),
        ],
    )

    doc.add_heading("9. Metricas utilizadas", level=1)
    doc.add_paragraph(
        "Para comparar modelos se usaron metricas comunes en deteccion de objetos. Estas metricas no significan solamente que el modelo "
        "acierta o falla, sino que consideran tanto si encontro los agaves como si las cajas quedaron bien ubicadas."
    )
    add_concept_table(
        doc,
        [
            ("Precision", "De todas las detecciones que hizo el modelo, indica que porcentaje fueron correctas. Alta precision significa pocos falsos positivos."),
            ("Recall", "De todos los agaves reales, indica que porcentaje encontro el modelo. Alto recall significa pocos objetos perdidos."),
            ("mAP50", "Promedio de precision considerando detecciones correctas cuando la caja tiene al menos 50% de solape con la etiqueta real."),
            ("mAP50-95", "Metrica mas estricta: evalua varios niveles de solape desde 50% hasta 95%. Exige cajas mejor ajustadas."),
            ("Best epoch", "Epoca donde el modelo obtuvo su mejor desempeno en validacion."),
            ("Checkpoint best.pt", "Archivo de pesos guardado en el mejor punto del entrenamiento."),
        ],
    )

    doc.add_heading("10. Benchmark de experimentos", level=1)
    doc.add_paragraph(
        "Se compararon 11 corridas. El modelo ganador fue yolo26s_sgd_cos_slow_120e con precision 0.693, recall 0.615, mAP50 0.654 "
        "y mAP50-95 0.213 en test."
    )
    add_benchmark_table(doc)
    add_caption(doc, "Tabla 1. Benchmark de corridas realizadas sobre el dataset de agave.")
    add_callout(
        doc,
        "Lectura del benchmark",
        "La corrida ganadora no fue simplemente la mas larga. Fue la que encontro mejor equilibrio entre precision, recall y mAP. "
        "Esto muestra que aumentar epocas o modificar la funcion de perdida no siempre mejora el resultado si el dataset aun es limitado.",
    )

    doc.add_heading("11. Interpretacion de resultados", level=1)
    add_bullets(
        doc,
        [
            "El modelo yolo26s tuvo mejor desempeno general que yolo26n porque tiene mayor capacidad para representar patrones visuales.",
            "El mejor punto estuvo en la epoca 56; despues, entrenar mas no produjo mejora suficiente.",
            "La precision 0.693 indica que, cuando el modelo detecta, una parte importante de sus detecciones son correctas, pero todavia hay falsos positivos.",
            "El recall 0.615 indica que aun puede perder algunos agaves, especialmente si estan pequenos, ocluidos o muy juntos.",
            "El mAP50 de 0.654 muestra que la deteccion general ya funciona como primera version.",
            "El mAP50-95 de 0.213 muestra que se debe mejorar el ajuste fino de las cajas, probablemente con mejores etiquetas y mas datos reales.",
        ],
    )

    doc.add_heading("12. Graficas y evidencia visual", level=1)
    doc.add_paragraph(
        "Las graficas se conservaron porque permiten explicar visualmente el proceso. Muestran como evolucionaron las perdidas y metricas "
        "durante el entrenamiento, la matriz de confusion y ejemplos de predicciones sobre validacion."
    )
    add_image_if_exists(
        doc,
        ROOT / "docs" / "results" / "best_yolo26s_results.png",
        "Figura 1. Curvas de entrenamiento y metricas del mejor modelo.",
        width=6.1,
    )
    doc.add_paragraph(
        "En estas curvas se observa la evolucion del aprendizaje. Lo importante es que las metricas de validacion permiten identificar "
        "si el modelo mejora de forma real o si solo esta ajustandose al conjunto de entrenamiento."
    )
    add_image_if_exists(
        doc,
        ROOT / "docs" / "results" / "best_yolo26s_confusion_matrix.png",
        "Figura 2. Matriz de confusion del mejor modelo.",
        width=4.6,
    )
    doc.add_paragraph(
        "La matriz de confusion ayuda a interpretar aciertos y errores. En este caso solo existe una clase, agave, por lo que el interes "
        "principal es observar detecciones correctas frente a omisiones o detecciones erroneas."
    )
    add_image_if_exists(
        doc,
        ROOT / "docs" / "results" / "best_yolo26s_val_predictions.jpg",
        "Figura 3. Predicciones en lote de validacion.",
        width=6.1,
    )
    doc.add_paragraph(
        "Las predicciones de validacion permiten revisar visualmente si las cajas caen sobre las plantas correctas. Esta revision visual es "
        "muy importante porque una metrica puede verse aceptable, pero la aplicacion real necesita cajas utiles en campo."
    )

    doc.add_heading("13. Simulacion sobre video", level=1)
    doc.add_paragraph(
        "La etapa mas reciente y mas importante para demostrar avance aplicado fue la simulacion sobre video. En lugar de evaluar solo imagenes "
        "individuales, se tomo un video real y se analizo un segmento especifico de 30 segundos: del minuto 1:30 al minuto 2:00."
    )
    doc.add_paragraph(
        "Esta prueba es relevante porque un video introduce problemas que no aparecen igual en imagenes sueltas: movimiento, cambios de escala, "
        "objetos que entran y salen del cuadro, agaves que se pierden algunos frames y luego vuelven a detectarse. Por eso se paso de simple "
        "deteccion a tracking y conteo."
    )
    add_key_value_table(
        doc,
        [
            ("Video fuente", r"C:\Users\iceve\Documents\3.mp4"),
            ("Segmento analizado", "Del minuto 1:30 al 2:00"),
            ("Duracion del segmento", "Aproximadamente 30 segundos"),
            ("Resolucion del video original", "2880 x 1620"),
            ("FPS aproximado", "15 FPS"),
            ("Archivo de segmento", "video_predictions/3_segment_01m30s_02m00s.mp4"),
            ("Modelo usado", r"runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt"),
            ("Parametros de inferencia", "conf=0.30, iou=0.40, imgsz=640, device=0"),
        ],
    )

    doc.add_heading("14. Que se hizo en la simulacion de video", level=1)
    add_numbered(
        doc,
        [
            "Se recorto el video original para trabajar solo con el tramo solicitado de 01:30 a 02:00.",
            "Se aplico el mejor modelo entrenado para detectar agaves en cada frame del segmento.",
            "Se uso tracking para asignar un ID a cada agave detectado a lo largo del video.",
            "Se dibujo una linea horizontal configurable en el frame para contar cruces.",
            "Se cambio la linea al 50% del frame, de acuerdo con la correccion solicitada.",
            "Se ajusto el conteo para considerar solamente IDs unicos que cruzan hacia abajo.",
            "Se genero un video de salida con cajas, IDs, trayectorias, linea de conteo y total acumulado.",
            "Se genero un CSV con los eventos de cruce para auditar el conteo.",
        ],
    )
    add_concept_table(
        doc,
        [
            ("Frame", "Cada imagen individual dentro de un video. Un video de 15 FPS contiene aproximadamente 15 frames por segundo."),
            ("ID de tracking", "Numero que el algoritmo asigna a un objeto para seguirlo a traves del tiempo."),
            ("Trayectoria", "Historial de posiciones de un objeto entre frames."),
            ("Linea al 50%", "Linea horizontal ubicada a la mitad vertical del frame. Se usa como referencia para contar cruces."),
            ("Cruce hacia abajo", "Evento donde el centro del objeto pasa de estar arriba de la linea a estar abajo de ella."),
            ("CSV de eventos", "Archivo tabular que guarda cada cruce contado, con frame, ID y datos de la deteccion."),
        ],
    )

    doc.add_heading("15. Correccion del conteo", level=1)
    doc.add_paragraph(
        "Durante la primera version del conteo, los IDs se visualizaban correctamente, pero el total se acumulaba mal. El problema no estaba "
        "principalmente en la deteccion, sino en la logica de conteo: un mismo objeto podia volver a generar eventos de conteo si se perdia "
        "por algunos frames o si la trayectoria cruzaba la linea de forma ambigua."
    )
    doc.add_paragraph(
        "La correccion consistio en crear un modo unique-id-only. Con este modo, cada ID de tracking se cuenta una sola vez por direccion. "
        "Asi, si el mismo agave continua apareciendo en el video, no vuelve a incrementar el total."
    )
    add_key_value_table(
        doc,
        [
            ("Script corregido", "scripts/track_count_line.py"),
            ("Modo usado", "--unique-id-only"),
            ("Direccion", "down / hacia abajo"),
            ("Linea", "50% del frame"),
            ("Tracker", "BoT-SORT mediante configs/tracking/botsort_agave.yaml"),
            ("Resultado corregido", "41 IDs unicos contados hacia abajo"),
            ("Duplicados suprimidos", "0 en modo unique-id-only, porque el control se hizo por ID unico"),
        ],
    )
    add_callout(
        doc,
        "Resultado clave de la simulacion",
        "Con la linea horizontal al 50% del frame y conteo solo hacia abajo, el sistema contabilizo 41 agaves unicos en el segmento de "
        "01:30 a 02:00. Este resultado ya representa una prueba funcional de deteccion + tracking + conteo.",
    )

    doc.add_heading("16. Interpretacion de la simulacion de video", level=1)
    add_bullets(
        doc,
        [
            "La simulacion demostro que el modelo puede usarse mas alla de imagenes estaticas.",
            "El tracking permitio mantener identidad de objetos mediante IDs, lo cual es esencial para conteo.",
            "La linea al 50% funciono como una regla simple y configurable para convertir detecciones en un evento medible.",
            "El conteo corregido de 41 IDs unicos muestra que la logica ya evita sumar repetidamente el mismo objeto cuando conserva el ID.",
            "La prueba todavia es una simulacion controlada; para uso real conviene comparar contra conteo manual y probar videos mas largos.",
        ],
    )
    add_key_value_table(
        doc,
        [
            ("Salida de video", "video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.mp4"),
            ("Salida CSV", "video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.csv"),
            ("IDs contados", "41 IDs unicos"),
            ("Lista de IDs", "1, 2, 6, 7, 8, 10, 22, 31, 32, 41, 49, 55, 59, 71, 84, 91, 108, 117, 123, 158, 159, 161, 168, 190, 191, 193, 203, 220, 224, 239, 278, 306, 309, 326, 335, 368, 397, 409, 451, 494, 513"),
        ],
    )

    doc.add_heading("17. Limitaciones actuales", level=1)
    add_bullets(
        doc,
        [
            "El dataset es pequeno para cubrir toda la variabilidad de campo: iluminacion, sombras, altura de camara, distancia y tipos de terreno.",
            "El modelo puede detectar dos veces un mismo agave si hay solapamiento, o si la caja se divide visualmente.",
            "El tracking actual funciona como primera etapa, pero si los IDs cambian durante videos largos, convendra integrar ReID o DeepSORT.",
            "La linea fija al 50% es util para prueba, pero podria necesitar ajuste segun perspectiva, velocidad de camara o zona de interes.",
            "El mAP50-95 todavia indica que la precision fina de las cajas puede mejorar.",
            "Hace falta una comparacion contra conteo manual para medir error absoluto del sistema.",
        ],
    )

    doc.add_heading("18. Recomendaciones para siguientes avances", level=1)
    add_numbered(
        doc,
        [
            "Capturar mas imagenes reales de agaves en condiciones variadas, antes de depender de aumentacion artificial.",
            "Revisar y limpiar etiquetas para mejorar calidad de cajas y subir mAP50-95.",
            "Probar yolo26m si la GPU lo permite, reduciendo batch si es necesario.",
            "Evaluar el modelo por escenarios: agaves juntos, agaves pequenos, sombras, movimiento y distancia.",
            "Crear una prueba formal de conteo manual vs automatico en videos completos.",
            "Probar tracking con DeepSORT o ReID si BoT-SORT empieza a perder identidades.",
            "Agregar zonas de interes configurables, no solo una linea horizontal.",
            "Exportar reportes automaticos por video: total contado, IDs, frames de cruce y capturas de eventos.",
        ],
    )

    doc.add_heading("19. Mensaje para presentacion de avances", level=1)
    doc.add_paragraph(
        "El avance no se limita a entrenar un modelo: se construyo un pipeline completo de trabajo. Ya existe un repositorio organizado, "
        "un entorno local reproducible con GPU, un dataset preparado, 11 experimentos comparados, un mejor modelo identificado y una prueba "
        "sobre video que produce un resultado medible."
    )
    add_bullets(
        doc,
        [
            "Mejor modelo actual: yolo26s_sgd_cos_slow_120e.",
            "Mejor metrica en test: mAP50 de 0.654.",
            "Conteo en video: 41 agaves unicos cruzando hacia abajo la linea al 50% del frame.",
            "Tecnica usada en video: deteccion YOLO + tracking BoT-SORT + conteo por cruce de linea.",
            "Siguiente objetivo: mejorar datos, validar conteo contra referencia manual y robustecer tracking.",
        ],
    )

    doc.add_heading("20. Rutas importantes", level=1)
    add_key_value_table(
        doc,
        [
            ("Mejor peso local", r"runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt"),
            ("Benchmark", "benchmark/agave_benchmark.csv"),
            ("Configuracion ganadora", "configs/training/historico_2026-05_06/yolo26s_sgd_cos_slow_120e.yaml"),
            ("Script de tracking", "scripts/track_count_line.py"),
            ("Tracker config", "configs/tracking/botsort_agave.yaml"),
            ("README tecnico", "README.md"),
            ("Documento de avances", "docs/reportes/Reporte_avances_deteccion_agave.docx"),
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
