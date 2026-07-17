from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Reporte_avance_tesis_deteccion_agave_yolo.docx"
RESULTS = ROOT / "docs" / "results"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(34, 34, 34)
MUTED = RGBColor(90, 90, 90)
FILL = "F2F4F7"
CALLOUT = "E8EEF5"
CODE_FILL = "F7F7F7"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": 80, "bottom": 80, "start": 120, "end": 120}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths: list[float], header=True) -> None:
    table.style = "Table Grid"
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(row.cells[i])
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
        for cell in table.rows[0].cells:
            shade(cell, FILL)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Reporte de avance de tesis")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = DARK_BLUE
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Deteccion de agave mediante redes neuronales convolucionales profundas YOLO26")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = INK
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Avance tecnico: preparacion de datos, entrenamiento, evaluacion, inferencia en video, tracking y conteo")
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED


def callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = ""
    style_table(table, [6.3], header=False)
    cell = table.rows[0].cells[0]
    shade(cell, CALLOUT)
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(text)
    doc.add_paragraph()


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def kv(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Elemento"
    table.rows[0].cells[1].text = "Detalle"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    style_table(table, [1.75, 4.55])


def concepts(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Explicacion"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    style_table(table, [1.55, 4.75])


def code_block(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    style_table(table, [6.3], header=False)
    cell = table.rows[0].cells[0]
    shade(cell, CODE_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(code)
    r.font.name = "Consolas"
    r.font.size = Pt(8)


def image(doc: Document, path: Path, caption: str, width=6.1) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def read_file(path: str, max_chars: int) -> str:
    text = (ROOT / path).read_text(encoding="utf-8", errors="replace")
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n...\n[Fragmento recortado. El archivo completo queda en el repositorio.]"


def benchmark_table(doc: Document) -> None:
    rows = []
    with (ROOT / "benchmark" / "agave_benchmark.csv").open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            clean = {(k or "").replace("\ufeff", "").strip().strip('"'): (v or "").strip().strip('"') for k, v in row.items()}
            rows.append(clean)
    table = doc.add_table(rows=1, cols=9)
    headers = ["Rank", "Corrida", "Modelo", "Ep.", "Best", "Prec.", "Recall", "mAP50", "mAP50-95"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        values = [
            row["rank"],
            row["run"],
            row["model"],
            row["epochs_run"],
            row["best_epoch"],
            row["test_precision"],
            row["test_recall"],
            row["test_mAP50"],
            row["test_mAP50_95"],
        ]
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
    style_table(table, [0.4, 1.55, 0.68, 0.42, 0.45, 0.52, 0.55, 0.55, 0.68])


def architecture_table(doc: Document) -> None:
    rows = [
        ("Backbone", "Conv, C3k2/C3k, SPPF, C2PSA", "Extrae caracteristicas visuales: bordes, texturas, forma radial de hojas y contexto del cultivo."),
        ("Neck", "Upsample, Concat, C3k2", "Fusiona caracteristicas de diferentes escalas para detectar agaves pequenos, medianos y completos."),
        ("Head", "Detect", "Predice cajas de deteccion, confianza y clase agave sobre multiples escalas."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Parte"
    table.rows[0].cells[1].text = "Bloques principales"
    table.rows[0].cells[2].text = "Funcion"
    for a, b, c in rows:
        cells = table.add_row().cells
        cells[0].text = a
        cells[1].text = b
        cells[2].text = c
    style_table(table, [1.0, 2.2, 3.1])


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    add_title(doc)

    callout(
        doc,
        "Resumen del avance",
        "Se desarrollo un flujo funcional de vision artificial para detectar agaves usando YOLO26: organizacion del dataset, configuracion del entorno con GPU, entrenamiento de modelos, evaluacion por metricas, seleccion del mejor checkpoint, inferencia en video, tracking con IDs y conteo por cruce de linea.",
    )

    doc.add_heading("1. Introduccion general del proyecto", level=1)
    doc.add_paragraph(
        "El presente avance se enfoca en la etapa de vision computacional de un sistema agricola orientado a cultivos de agave. El objetivo es construir y validar una base de deteccion automatica capaz de reconocer plantas de agave en imagenes RGB y video, usando redes neuronales convolucionales profundas."
    )
    doc.add_paragraph(
        "Este trabajo se relaciona con agricultura de precision, ya que la deteccion visual puede convertirse posteriormente en informacion util para monitoreo, estimacion de presencia de plantas, conteo, analisis de surcos o activacion futura de mecanismos agricolas selectivos. Aunque la tesis contempla avanzar hacia deteccion de maleza y posible implementacion embebida, en este avance se consolido primero la deteccion de la clase agave y la prueba aplicada sobre video."
    )
    concepts(
        doc,
        [
            ("Vision artificial", "Campo que permite a una computadora interpretar imagenes o video para reconocer objetos, patrones o eventos."),
            ("Deteccion de objetos", "Tarea donde el modelo identifica donde esta un objeto y dibuja una caja alrededor de el."),
            ("YOLO26", "Modelo de deteccion usado como base para entrenar un detector especifico de agave."),
            ("Red neuronal convolucional", "Arquitectura especializada en imagenes que aprende filtros visuales como bordes, texturas y formas."),
        ],
    )

    doc.add_heading("2. Metodologia desarrollada", level=1)
    doc.add_paragraph(
        "La metodologia seguida conserva la estructura del protocolo de tesis: adquisicion y preparacion de datos, diseno y entrenamiento del modelo, evaluacion, prueba sobre video y preparacion para etapas posteriores de implementacion."
    )
    numbered(
        doc,
        [
            "Organizar el proyecto como repositorio Git y conectarlo a GitHub.",
            "Preparar el dataset en formato YOLO26 con train, validation y test.",
            "Crear un entorno local con Pipenv y librerias de entrenamiento e inferencia.",
            "Validar que PyTorch pudiera usar la GPU NVIDIA mediante CUDA.",
            "Entrenar multiples variantes YOLO26 con diferentes parametros.",
            "Comparar resultados mediante precision, recall, mAP50 y mAP50-95.",
            "Seleccionar el mejor modelo.",
            "Aplicar el modelo a un segmento de video y generar resultados visuales.",
            "Implementar tracking y conteo de agaves mediante IDs unicos y linea horizontal.",
        ],
    )

    doc.add_heading("2.1 Adquisicion y preparacion de datos", level=2)
    doc.add_paragraph(
        "El dataset fue exportado desde Roboflow y organizado localmente en formato YOLO. Para este avance se trabajo con una clase: agave. Las imagenes fueron divididas en subconjuntos de entrenamiento, validacion y prueba para poder entrenar y evaluar de manera ordenada."
    )
    kv(
        doc,
        [
            ("Ruta del dataset", "datasets/yolo26"),
            ("Clase", "agave"),
            ("Imagenes totales", "404"),
            ("Train", "295 imagenes"),
            ("Validation", "73 imagenes"),
            ("Test", "36 imagenes"),
            ("Archivo de configuracion", "datasets/yolo26/data.yaml"),
        ],
    )
    code_block(doc, read_file("datasets/yolo26/data.yaml", 900))

    doc.add_heading("2.2 Configuracion del entorno de trabajo", level=2)
    doc.add_paragraph(
        "El entorno se preparo usando Pipenv, no Conda. Esto permite conservar las dependencias del proyecto en Pipfile y Pipfile.lock. La parte clave fue instalar PyTorch con soporte CUDA, para que el entrenamiento pudiera ejecutarse en la GPU NVIDIA del equipo."
    )
    kv(
        doc,
        [
            ("Python", "3.12.9"),
            ("GPU", "NVIDIA GeForce RTX 5050 Laptop GPU, 8 GB"),
            ("PyTorch", "2.10.0+cu128"),
            ("Torchvision", "0.25.0+cu128"),
            ("Framework", "Ultralytics / YOLO26"),
            ("Servicios/librerias adicionales", "Roboflow, inference-sdk, python-dotenv"),
        ],
    )
    code_block(doc, read_file("Pipfile", 1800))

    doc.add_heading("2.3 Diseno, entrenamiento y validacion del modelo", level=2)
    doc.add_paragraph(
        "Se entrenaron modelos YOLO26 partiendo de pesos preentrenados. Se probaron variantes yolo26n y yolo26s. La version n es mas ligera y rapida; la version s tiene mayor capacidad de aprendizaje y fue la que finalmente obtuvo el mejor desempeno."
    )
    concepts(
        doc,
        [
            ("Epoca", "Una pasada completa por las imagenes de entrenamiento."),
            ("Batch", "Numero de imagenes procesadas antes de actualizar los pesos del modelo."),
            ("Learning rate", "Tamaño de los pasos de aprendizaje durante la optimizacion."),
            ("Optimizador", "Algoritmo que actualiza pesos; se exploraron SGD y AdamW."),
            ("Data augmentation", "Transformaciones aplicadas para simular variabilidad en los datos."),
            ("Early stopping", "Detencion automatica cuando el modelo deja de mejorar."),
        ],
    )

    doc.add_heading("3. Cronograma y avance real", level=1)
    doc.add_paragraph(
        "El avance realizado corresponde principalmente a la fase de vision computacional del protocolo. Se alcanzo una base funcional que permite entrenar, evaluar, inferir en video y contar objetos. La implementacion embebida en Raspberry Pi y la comunicacion con actuadores quedan como siguientes fases."
    )
    kv(
        doc,
        [
            ("Conceptualizacion", "Se definio el objetivo de detectar agave como primera clase de trabajo."),
            ("Base de datos", "Se organizo un dataset YOLO26 con 404 imagenes etiquetadas."),
            ("Entorno", "Se instalo y valido un entorno local con GPU para entrenamiento."),
            ("Modelos", "Se entrenaron 11 corridas comparativas."),
            ("Video", "Se probo el mejor modelo en un segmento real de 30 segundos."),
            ("Tracking", "Se agrego seguimiento por ID y conteo por linea."),
            ("Pendiente", "Validacion contra conteo manual, deteccion de maleza e implementacion embebida."),
        ],
    )

    doc.add_heading("4. Entrenamientos realizados", level=1)
    doc.add_paragraph(
        "Durante el proceso se realizaron diversas corridas con cambios en epocas, optimizador, learning rate, funcion de perdida y aumentacion. El objetivo no fue solo obtener un resultado, sino explorar que combinacion ofrecía mejor equilibrio entre precision, recall y mAP."
    )
    benchmark_table(doc)
    callout(
        doc,
        "Mejor modelo obtenido",
        "La mejor corrida fue yolo26s_sgd_cos_slow_120e, con precision 0.693, recall 0.615, mAP50 0.654 y mAP50-95 0.213 en test. El peso final usado fue runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt.",
    )

    doc.add_heading("4.1 Parametros del mejor entrenamiento", level=2)
    kv(
        doc,
        [
            ("Modelo base", "yolo26s.pt"),
            ("Epocas configuradas", "120"),
            ("Epocas corridas", "86"),
            ("Mejor epoca", "56 en benchmark final"),
            ("Batch", "8"),
            ("Tamano de imagen", "640"),
            ("Optimizador", "SGD"),
            ("Learning rate inicial", "0.005"),
            ("Scheduler", "Cosine LR"),
            ("Patience", "30"),
            ("Mosaic", "0.0"),
            ("Rotacion", "0 grados"),
            ("Translacion", "0.0"),
            ("Flip vertical", "0.5"),
            ("Flip horizontal", "0.5"),
            ("Loss", "box=7.5, cls=0.5, dfl=1.5"),
        ],
    )
    code_block(doc, read_file("training_configs/yolo26s_sgd_cos_slow_120e.yaml", 4200))

    doc.add_heading("5. Resultados y evidencia visual", level=1)
    doc.add_paragraph(
        "Las graficas permiten analizar el comportamiento del entrenamiento y explicar si el modelo realmente aprendio. Se incluyen curvas de entrenamiento, una metrica operacional tipo accuracy/F1, matriz de confusion legible y ejemplos de prediccion."
    )
    image(doc, RESULTS / "best_yolo26s_results.png", "Figura 1. Curvas originales del entrenamiento del mejor modelo.", 6.1)
    image(doc, RESULTS / "best_yolo26s_operational_accuracy.png", "Figura 2. ACC operacional, precision, recall y mAP50 durante entrenamiento.", 6.1)
    image(doc, RESULTS / "best_yolo26s_confusion_matrix_readable.png", "Figura 3. Matriz de confusion redibujada con mayor contraste.", 5.5)
    image(doc, RESULTS / "best_yolo26s_val_predictions.jpg", "Figura 4. Predicciones visuales sobre imagenes de validacion.", 6.1)
    concepts(
        doc,
        [
            ("Precision", "Mide que proporcion de detecciones marcadas como agave fueron correctas."),
            ("Recall", "Mide que proporcion de agaves reales fueron encontrados."),
            ("mAP50", "Evalua detecciones correctas con al menos 50% de solape entre caja predicha y caja real."),
            ("mAP50-95", "Evalua cajas con criterios mas estrictos; por eso suele ser menor."),
            ("ACC operacional", "Metrica agregada basada en F1 para explicar balance entre precision y recall. No sustituye mAP."),
        ],
    )

    doc.add_heading("6. Arquitectura del modelo", level=1)
    doc.add_paragraph(
        "A grandes rasgos, YOLO26 opera como una red de deteccion formada por tres partes: backbone, neck y head. El checkpoint final tiene 24 bloques principales y 9,948,638 parametros. La arquitectura aprende progresivamente desde rasgos simples, como bordes y texturas, hasta rasgos de mayor nivel, como forma general y ubicacion del agave."
    )
    architecture_table(doc)
    concepts(
        doc,
        [
            ("Conv2D", "Capa convolucional que extrae patrones locales de la imagen."),
            ("BatchNorm", "Normaliza activaciones y estabiliza entrenamiento."),
            ("SiLU", "Funcion de activacion que permite aprender relaciones no lineales."),
            ("C3k2/C3k", "Bloques compuestos que reutilizan caracteristicas y mejoran aprendizaje."),
            ("SPPF", "Pooling piramidal rapido que captura contexto a varias escalas."),
            ("C2PSA", "Bloque de atencion que ayuda a priorizar informacion relevante."),
            ("Upsample", "Aumenta resolucion de mapas de caracteristicas."),
            ("Concat", "Une caracteristicas de diferentes niveles."),
            ("Detect", "Cabeza final que predice cajas, confianza y clase."),
        ],
    )

    doc.add_heading("7. Inferencia en video, tracking y conteo", level=1)
    doc.add_paragraph(
        "La etapa aplicada consistio en usar el mejor modelo sobre un video real. Se proceso el segmento de 01:30 a 02:00 del video 3.mp4. Posteriormente se agrego tracking para asignar un ID a cada agave y contar solo aquellos que cruzaran una linea horizontal ubicada al 50% del frame."
    )
    kv(
        doc,
        [
            ("Video fuente", r"C:\Users\iceve\Documents\3.mp4"),
            ("Segmento", "01:30 a 02:00"),
            ("Frames procesados", "450 aproximadamente"),
            ("Modelo", "yolo26s_sgd_cos_slow_120e/weights/best.pt"),
            ("Confianza", "0.30"),
            ("IoU", "0.40"),
            ("Tracker", "BoT-SORT"),
            ("Linea de conteo", "50% del frame"),
            ("Direccion", "Hacia abajo"),
            ("Resultado", "41 agaves unicos contados"),
        ],
    )
    image(doc, RESULTS / "final_tracking_video_frame.jpg", "Figura 5. Captura del video final con detecciones, IDs, linea y conteo.", 6.1)
    callout(
        doc,
        "Resultado aplicado",
        "El sistema logro detectar agaves en video, asignar IDs de seguimiento y contabilizar 41 agaves unicos cruzando hacia abajo la linea al 50% del frame. Esta es la evidencia mas cercana a una aplicacion en campo dentro del avance actual.",
    )

    doc.add_heading("7.1 Codigo usado para tracking", level=2)
    doc.add_paragraph(
        "El script track_count_line.py implementa lectura de video, inferencia YOLO, tracking persistente, dibujo de cajas e IDs, conteo de cruce y exportacion de CSV. El modo --unique-id-only evita contar dos veces el mismo ID."
    )
    code_block(
        doc,
        "$env:PIPENV_VENV_IN_PROJECT='1'\npy -3.12 -m pipenv run python scripts/track_count_line.py `\n  --model \"runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt\" `\n  --source \"video_predictions/3_segment_01m30s_02m00s.mp4\" `\n  --output \"video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.mp4\" `\n  --csv \"video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.csv\" `\n  --tracker \"tracking_configs/botsort_agave.yaml\" `\n  --line-ratio 0.50 `\n  --direction down `\n  --conf 0.30 `\n  --iou 0.40 `\n  --imgsz 640 `\n  --device 0 `\n  --unique-id-only `\n  --show-conf",
    )
    code_block(doc, read_file("tracking_configs/botsort_agave.yaml", 1200))

    doc.add_heading("8. Comparacion entre avance real y siguientes fases", level=1)
    doc.add_paragraph(
        "El avance cubre de forma satisfactoria la fase de vision computacional para la clase agave. Las etapas siguientes deben enfocarse en ampliar el dataset, incluir maleza como clase adicional, validar conteo contra referencia manual y preparar exportacion del modelo para una plataforma embebida."
    )
    bullets(
        doc,
        [
            "Agregar imagenes reales de diferentes parcelas, alturas, iluminacion y angulos.",
            "Incluir etiquetas de maleza y revisar consistencia de anotaciones.",
            "Probar modelos mas grandes o versiones optimizadas para inferencia embebida.",
            "Comparar conteo automatico contra conteo manual para estimar error.",
            "Exportar el modelo a ONNX, TensorRT u otro formato segun la plataforma final.",
            "Evaluar FPS y latencia en Raspberry Pi u otro dispositivo embebido.",
            "Definir como las detecciones se traduciran en senales para actuadores.",
        ],
    )

    doc.add_heading("9. Conclusion del avance", level=1)
    doc.add_paragraph(
        "El trabajo realizado demuestra que ya existe una base funcional para deteccion de agave con redes neuronales profundas. Se logro pasar de un conjunto de imagenes etiquetadas a un sistema capaz de entrenar, evaluar, inferir en video, seguir objetos y contar cruces. El mejor modelo actual, yolo26s_sgd_cos_slow_120e, ofrece un punto de partida solido para continuar hacia deteccion de maleza, validacion en campo e implementacion embebida."
    )
    doc.add_paragraph(
        "Para la tesis, este avance puede presentarse como una primera etapa experimental completa: construccion del dataset, seleccion de arquitectura, entrenamiento comparativo, analisis de metricas y prueba aplicada sobre video. El siguiente paso metodologico es ampliar la base de datos y validar el sistema en condiciones mas cercanas a una operacion real."
    )

    doc.add_heading("10. Archivos principales generados", level=1)
    kv(
        doc,
        [
            ("Mejor modelo", r"runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt"),
            ("Benchmark", "benchmark/agave_benchmark.csv"),
            ("Configuracion ganadora", "training_configs/yolo26s_sgd_cos_slow_120e.yaml"),
            ("Script tracking", "scripts/track_count_line.py"),
            ("Config tracker", "tracking_configs/botsort_agave.yaml"),
            ("Video final", "video_predictions/3_segment_01m30s_02m00s_tracking_line50_down_unique.mp4"),
            ("Reporte", "docs/Reporte_avance_tesis_deteccion_agave_yolo.docx"),
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
