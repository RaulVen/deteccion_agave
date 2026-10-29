from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reportes" / "Reporte_especifico_modelo_agave.docx"
RESULTS = ROOT / "docs" / "results"
RUN = ROOT / "runs" / "detect" / "agave_experiments" / "yolo26s_sgd_cos_slow_120e"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(34, 34, 34)
MUTED = RGBColor(90, 90, 90)
HEADER_FILL = "F2F4F7"
CALLOUT_FILL = "E8EEF5"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": 80, "bottom": 80, "start": 120, "end": 120}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths: list[float], header=True) -> None:
    table.style = "Table Grid"
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(row.cells[i])
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
        for cell in table.rows[0].cells:
            shade(cell, HEADER_FILL)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9)
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def add_kv(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Elemento"
    table.rows[0].cells[1].text = "Detalle"
    for k, v in rows:
        row = table.add_row()
        row.cells[0].text = k
        row.cells[1].text = v
    style_table(table, [1.8, 4.5])


def add_concepts(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concepto"
    table.rows[0].cells[1].text = "Explicacion"
    for k, v in rows:
        row = table.add_row()
        row.cells[0].text = k
        row.cells[1].text = v
    style_table(table, [1.55, 4.75])


def add_callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    row = table.rows[0]
    row.cells[0].text = ""
    style_table(table, [6.3], header=False)
    cell = row.cells[0]
    shade(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(text)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def image(doc: Document, path: Path, text: str, width=6.1) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    caption(doc, text)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color


def architecture_table(doc: Document) -> None:
    rows = [
        ("0", "Conv", "3 -> 32", "Primera extraccion de bordes, texturas y contrastes desde la imagen RGB."),
        ("1", "Conv", "32 -> 64", "Reduce resolucion y aumenta canales para aprender patrones mas ricos."),
        ("2", "C3k2", "64 -> 128", "Bloque tipo CSP/bottleneck: conserva informacion y extrae rasgos sin elevar demasiado el costo."),
        ("3", "Conv", "128 -> 128", "Downsampling: reduce tamano espacial para captar contexto mas amplio."),
        ("4", "C3k2", "128 -> 256", "Aprende partes de agave: hojas, bordes radiales y texturas repetidas."),
        ("5", "Conv", "256 -> 256", "Nuevo downsampling para pasar a rasgos mas abstractos."),
        ("6", "C3k2/C3k", "256 -> 256", "Bloque profundo con bottlenecks; mezcla detalles locales y contexto."),
        ("7", "Conv", "256 -> 512", "Aumenta canales para representar objetos completos y patrones complejos."),
        ("8", "C3k2/C3k", "512 -> 512", "Backbone profundo: consolida rasgos visuales del agave."),
        ("9", "SPPF", "512 -> 512", "Pooling piramidal rapido; captura contexto a multiples escalas."),
        ("10", "C2PSA", "512 -> 512", "Bloque de atencion; ayuda a priorizar zonas relevantes de la imagen."),
        ("11", "Upsample", "-", "Aumenta resolucion de mapas para recuperar detalle espacial."),
        ("12", "Concat", "-", "Fusiona mapas profundos con mapas de capas anteriores."),
        ("13", "C3k2", "768 -> 256", "Refina la fusion para deteccion de objetos medianos."),
        ("14", "Upsample", "-", "Vuelve a subir resolucion para objetos mas pequenos o detalles finos."),
        ("15", "Concat", "-", "Combina detalle fino del backbone con informacion semantica."),
        ("16", "C3k2", "512 -> 128", "Refina caracteristicas para escala fina."),
        ("17", "Conv", "128 -> 128", "Downsampling en el neck para reconstruir piramide multi-escala."),
        ("18", "Concat", "-", "Fusiona escala fina y media."),
        ("19", "C3k2", "384 -> 256", "Refina escala media para deteccion estable."),
        ("20", "Conv", "256 -> 256", "Downsampling hacia escala grande."),
        ("21", "Concat", "-", "Fusiona informacion de alto nivel."),
        ("22", "C3k2", "768 -> 512", "Refina escala grande para objetos completos."),
        ("23", "Detect", "3 escalas -> agave", "Cabeza de deteccion: predice cajas, confianza y clase agave."),
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["Bloque", "Tipo", "Canales/escala", "Funcion en el modelo"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for data in rows:
        row = table.add_row()
        for i, value in enumerate(data):
            row.cells[i].text = value
    style_table(table, [0.55, 1.0, 1.15, 3.6])


def params_table(doc: Document) -> None:
    path = ROOT / "docs" / "model_layers_summary.tsv"
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Bloque"
    table.rows[0].cells[1].text = "Tipo"
    table.rows[0].cells[2].text = "Parametros"
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            idx, layer, params = line.split("\t")
            row = table.add_row()
            row.cells[0].text = idx
            row.cells[1].text = layer
            row.cells[2].text = f"{int(params):,}"
    style_table(table, [0.75, 2.1, 1.6])
    for row in table.rows[1:]:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Reporte especifico: metricas visuales y arquitectura del modelo YOLO26")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = DARK_BLUE
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Proyecto de deteccion de agave")
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED

    add_callout(
        doc,
        "Objetivo de este reporte",
        "Complementar el reporte de avances con tres puntos especificos: matriz de confusion legible, grafica de ACC operacional y explicacion detallada de la arquitectura YOLO26 usada.",
    )

    doc.add_heading("1. Modelo utilizado", level=1)
    add_kv(
        doc,
        [
            ("Mejor modelo", "yolo26s_sgd_cos_slow_120e"),
            ("Peso utilizado", r"runs/detect/agave_experiments/yolo26s_sgd_cos_slow_120e/weights/best.pt"),
            ("Modelo base", "yolo26s.pt"),
            ("Clase entrenada", "agave"),
            ("Parametros totales", "9,948,638"),
            ("Bloques principales", "24 bloques secuenciales, incluyendo backbone, neck y cabeza Detect"),
            ("Metricas test del benchmark", "Precision 0.693, Recall 0.615, mAP50 0.654, mAP50-95 0.213"),
        ],
    )

    doc.add_heading("2. Aclaracion sobre ACC en deteccion de objetos", level=1)
    doc.add_paragraph(
        "En clasificacion tradicional se habla de accuracy como porcentaje de ejemplos clasificados correctamente. En deteccion de objetos "
        "esto no es tan directo, porque el modelo debe hacer dos cosas al mismo tiempo: encontrar la ubicacion del objeto y clasificarlo. "
        "Por eso YOLO reporta principalmente precision, recall, mAP50 y mAP50-95."
    )
    add_callout(
        doc,
        "ACC operacional agregada",
        "Para satisfacer la lectura de accuracy en graficas, se agrego una metrica operacional basada en F1: 2 * precision * recall / (precision + recall). "
        "Esta metrica resume el balance entre falsos positivos y objetos no detectados. No sustituye mAP, pero ayuda a explicar desempeno de forma sencilla.",
    )
    add_concepts(
        doc,
        [
            ("Precision", "De todo lo que el modelo marco como agave, que proporcion realmente fue correcta."),
            ("Recall", "De todos los agaves reales, que proporcion logro encontrar el modelo."),
            ("ACC operacional / F1", "Resumen balanceado de precision y recall. Sube solo cuando el modelo detecta bien y pierde pocos objetos."),
            ("mAP50", "Evalua detecciones correctas con al menos 50% de solape entre caja predicha y caja real."),
            ("mAP50-95", "Evalua cajas con criterios mas estrictos; es mas dificil y refleja mejor la calidad fina de localizacion."),
        ],
    )

    doc.add_heading("3. Graficas y evidencia visual conservadas", level=1)
    doc.add_paragraph(
        "Se conservan las graficas originales del entrenamiento y las predicciones visuales. Adicionalmente se agregan dos nuevas figuras: "
        "una grafica de ACC operacional y una matriz de confusion redibujada con contraste alto."
    )
    image(doc, RESULTS / "best_yolo26s_results.png", "Figura 1. Curvas originales de entrenamiento del mejor modelo.", width=6.1)
    image(doc, RESULTS / "best_yolo26s_operational_accuracy.png", "Figura 2. Nueva grafica: ACC operacional junto con precision, recall y mAP50.", width=6.1)
    image(doc, RESULTS / "best_yolo26s_val_predictions.jpg", "Figura 3. Evidencia visual: predicciones sobre validacion.", width=6.1)

    doc.add_heading("4. Matriz de confusion corregida visualmente", level=1)
    doc.add_paragraph(
        "La matriz original generada por YOLO puede ser dificil de leer cuando las celdas relacionadas con background quedan claras o casi blancas. "
        "Para la presentacion se agrego una version redibujada con contraste alto, texto grande y explicacion de cada celda."
    )
    image(doc, RESULTS / "best_yolo26s_confusion_matrix.png", "Figura 4. Matriz original generada por YOLO.", width=4.7)
    image(doc, RESULTS / "best_yolo26s_confusion_matrix_readable.png", "Figura 5. Matriz redibujada con contraste alto y etiquetas explicativas.", width=5.6)
    add_concepts(
        doc,
        [
            ("TP - verdaderos positivos", "Agaves reales que el modelo detecto correctamente."),
            ("FN - falsos negativos", "Agaves reales que el modelo no detecto."),
            ("FP - falsos positivos", "Detecciones que el modelo marco como agave, pero no correspondian a un agave real."),
            ("Background", "Fondo de la imagen. En deteccion de objetos no se cuenta como una clase positiva normal; sirve para entender errores contra el fondo."),
            ("TN - verdaderos negativos", "En deteccion no se interpreta igual que en clasificacion, porque hay infinitas regiones de fondo posibles. Por eso se marca como no aplicable."),
        ],
    )

    doc.add_heading("5. Arquitectura general del modelo", level=1)
    doc.add_paragraph(
        "El modelo YOLO26 usado sigue una organizacion tipica de detectores modernos: backbone, neck y head. El backbone extrae caracteristicas "
        "de la imagen; el neck fusiona informacion de distintas escalas; y la cabeza Detect produce las cajas finales, la confianza y la clase agave."
    )
    add_concepts(
        doc,
        [
            ("Backbone", "Parte inicial que transforma pixeles en mapas de caracteristicas. Aprende bordes, texturas, formas y patrones del agave."),
            ("Neck", "Parte intermedia que combina mapas de caracteristicas de diferentes escalas para detectar objetos pequenos, medianos y grandes."),
            ("Head / Detect", "Parte final que convierte caracteristicas en predicciones: ubicacion de caja, confianza y clase."),
            ("Feature map", "Representacion interna de la imagen despues de pasar por capas convolucionales."),
            ("Multi-escala", "Estrategia para detectar objetos de diferentes tamanos usando mapas con resoluciones distintas."),
        ],
    )

    doc.add_heading("6. Capas y bloques principales", level=1)
    doc.add_paragraph(
        "La arquitectura real del checkpoint se resume en 24 bloques principales. Internamente, cada bloque puede contener varias operaciones, "
        "por ejemplo convolucion, normalizacion BatchNorm y activacion SiLU."
    )
    architecture_table(doc)

    doc.add_heading("7. Para que sirve cada tipo de capa", level=1)
    add_concepts(
        doc,
        [
            ("Conv2D", "Aplica filtros sobre la imagen o feature map. Detecta patrones locales como bordes, lineas, texturas y formas."),
            ("BatchNorm2D", "Normaliza activaciones para que el entrenamiento sea mas estable y rapido."),
            ("SiLU", "Funcion de activacion no lineal. Permite que la red aprenda relaciones complejas, no solo transformaciones lineales."),
            ("Bottleneck", "Sub-bloque que reduce y expande canales para aprender de forma eficiente sin disparar parametros."),
            ("C3k2 / C3k", "Bloques compuestos tipo CSP que dividen y combinan rutas internas para mejorar aprendizaje y flujo de gradiente."),
            ("SPPF", "Spatial Pyramid Pooling Fast. Resume contexto en varias escalas sin perder demasiada eficiencia."),
            ("C2PSA", "Bloque de atencion que ayuda a enfatizar regiones o canales mas informativos."),
            ("Upsample", "Aumenta resolucion de mapas para recuperar detalle espacial."),
            ("Concat", "Une mapas de caracteristicas de diferentes etapas para mezclar detalle fino y contexto profundo."),
            ("Detect", "Capa final que predice cajas, confianza y clase agave en varias escalas."),
        ],
    )

    doc.add_heading("8. Conteo de parametros por bloque", level=1)
    doc.add_paragraph(
        "El conteo de parametros indica cuanta capacidad de aprendizaje tiene cada bloque. Los bloques profundos y la cabeza de deteccion "
        "concentran una parte importante de los parametros porque trabajan con representaciones mas complejas."
    )
    params_table(doc)

    doc.add_heading("9. Relacion entre arquitectura y deteccion de agave", level=1)
    add_bullets(
        doc,
        [
            "Las primeras convoluciones aprenden rasgos simples: bordes de hojas, cambios de color y textura del suelo.",
            "Los bloques C3k2 combinan rasgos para formar patrones mas especificos: roseta del agave, hojas alargadas y repeticion radial.",
            "SPPF permite que el modelo vea contexto: diferencia entre una planta completa y fragmentos del fondo.",
            "El neck con Upsample y Concat mezcla detalle fino con contexto profundo, util cuando hay agaves de tamanos diferentes.",
            "La cabeza Detect genera las cajas finales y decide si cada region corresponde a la clase agave.",
        ],
    )

    doc.add_heading("10. Conclusiones especificas", level=1)
    add_bullets(
        doc,
        [
            "Se corrigio el problema visual de la matriz de confusion agregando una version con contraste alto.",
            "Se conservo la evidencia visual previa y se agrego una grafica de ACC operacional.",
            "Se explico que ACC clasica no es la metrica principal en deteccion; se usa una aproximacion basada en F1 para presentacion.",
            "La arquitectura del modelo tiene 24 bloques principales y 9,948,638 parametros.",
            "El modelo combina convoluciones, bloques CSP, pooling piramidal, atencion, fusion multi-escala y cabeza de deteccion.",
        ],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
